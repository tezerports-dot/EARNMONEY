"""Bank-upload exports."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from app import db, export
from app.timeutil import current_month
from tests.conftest import join_both


def _prepare(world) -> str:
    month = current_month()
    join_both(5001)
    join_both(5002)
    db.record_activity(5001)
    db.record_activity(5002)
    db.save_bank_details(5001, "Alice Example", "0012345678901234", "HDFC0001234", "alice@okhdfc")
    db.create_withdrawal(5001, month, 10.0)
    return month


def test_csv_has_the_bank_columns(world):
    month = _prepare(world)
    payload, skipped = export.to_csv(month)
    assert skipped == []

    rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
    assert rows[0] == export.COLUMNS
    assert rows[1] == ["Alice Example", "0012345678901234", "HDFC0001234", "alice@okhdfc", "10.00"]


def test_csv_uses_crlf_and_a_bom_for_bank_portals(world):
    month = _prepare(world)
    payload, _ = export.to_csv(month)
    assert payload.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in payload


def test_xlsx_keeps_account_numbers_as_text(world):
    month = _prepare(world)
    payload, _ = export.to_xlsx(month)
    sheet = load_workbook(io.BytesIO(payload)).active

    assert [cell.value for cell in sheet[1]] == export.COLUMNS
    account_cell = sheet.cell(row=2, column=2)
    assert account_cell.value == "0012345678901234"  # leading zero survives
    assert account_cell.number_format == "@"
    assert sheet.cell(row=2, column=5).value == 10.0


def test_requests_without_bank_details_are_skipped_and_reported(world):
    month = _prepare(world)
    # Bob requests a payout but never submitted bank details.
    db.create_withdrawal(5002, month, 10.0)

    rows, skipped = export.collect_rows(month)
    assert len(rows) == 1
    assert skipped == [db.get_user(5002)["uid"]]


def test_rejected_and_paid_rows_stay_out_of_the_default_export(world):
    month = _prepare(world)
    request = db.get_withdrawal(5001, month)
    db.set_withdrawal_status(int(request["id"]), "rejected")

    rows, _ = export.collect_rows(month)
    assert rows == []

    db.set_withdrawal_status(int(request["id"]), "approved")
    rows, _ = export.collect_rows(month)
    assert len(rows) == 1
